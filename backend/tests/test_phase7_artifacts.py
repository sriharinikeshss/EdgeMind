"""
Phase 7 tests — Artifact Generation (DOCX/XLSX/PDF/CSV/JSON).

Covers:
  - Each writer produces a real, re-openable file with a correct SHA-256 hash.
  - Tools are registered and callable through the ToolRegistry.
  - Planner detects deliverable requests and appends a generate_* step when
    the model's own plan omits one.
  - Full agent-loop integration: a report request produces a persisted
    Artifact row, a downloadable file, and a working integrity check.
  - Tamper/missing-file edge cases on the download & verify endpoints.
"""
import json
import os
import shutil
import sys

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from main import app
from database.session import Base, engine, get_db
from database.models import Artifact, Task, TaskStep, ToolCall, ModelRoute
from api.auth import get_current_user, UserInfo
from sqlalchemy.orm import sessionmaker

from artifacts.docx_writer import DocxWriter
from artifacts.xlsx_writer import XlsxWriter
from artifacts.pdf_writer import PdfWriter
from artifacts.data_writer import CsvWriter, JsonWriter
from artifacts.storage import ARTIFACTS_DIR, resolve_artifact_path, sha256_file
from orchestrator.planner import Planner

TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


def override_get_current_user():
    return UserInfo(username="test-operator", role="operator")


app.dependency_overrides[get_db] = override_get_db
app.dependency_overrides[get_current_user] = override_get_current_user
client = TestClient(app)


TEST_TASK_ID = "phase7-test"


@pytest.fixture(autouse=True)
def setup_db():
    app.dependency_overrides.clear()
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = override_get_current_user
    Base.metadata.create_all(bind=engine)

    db = TestingSessionLocal()
    # artifacts.task_id is FK-constrained to tasks.id — writer unit tests use a
    # fixed task_id, so a matching Task row must exist for those inserts to succeed.
    if not db.query(Task).filter(Task.id == TEST_TASK_ID).first():
        db.add(Task(id=TEST_TASK_ID, description="Phase 7 test fixture task", status="COMPLETED"))
        db.commit()
    db.close()

    yield

    # Child tables must be cleared before their parent Task rows to satisfy FKs.
    db = TestingSessionLocal()
    db.query(Artifact).delete()
    db.query(TaskStep).delete()
    db.query(ToolCall).delete()
    db.query(ModelRoute).delete()
    db.query(Task).delete()
    db.commit()
    db.close()
    app.dependency_overrides.clear()


@pytest.fixture(autouse=True, scope="module")
def cleanup_artifact_dir():
    yield
    shutil.rmtree(os.path.join(ARTIFACTS_DIR, "phase7-test"), ignore_errors=True)


# ── Writers — unit level ─────────────────────────────────────────────────────

def test_docx_writer_creates_real_file_and_hash_matches():
    result = DocxWriter().create_docx(
        task_id="phase7-test",
        title="Inspection Summary",
        content="First paragraph.\n\nSecond paragraph with findings.",
        table={"headers": ["Item", "Status"], "rows": [["Valve A", "OK"], ["Valve B", "Leak"]]},
        citations=[{"doc_id": "sop-1", "text": "Pressure limit is 150 psi."}],
    )
    assert result["status"] == "ok"
    assert result["filename"].endswith(".docx")
    assert os.path.isfile(result["path"])
    assert result["file_hash"] == sha256_file(result["path"])

    # Must be a real, re-openable docx (python-docx can parse it back)
    from docx import Document
    doc = Document(result["path"])
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "First paragraph." in full_text
    assert len(doc.tables) == 1


def test_xlsx_writer_creates_real_file_with_correct_cells():
    result = XlsxWriter().create_xlsx(
        task_id="phase7-test",
        title="Readings",
        headers=["Sensor", "Value"],
        rows=[["PT-101", 42], ["PT-102", 37]],
        citations=["Manual v2, page 4"],
    )
    assert result["status"] == "ok"
    assert result["filename"].endswith(".xlsx")
    assert result["file_hash"] == sha256_file(result["path"])

    from openpyxl import load_workbook
    wb = load_workbook(result["path"])
    ws = wb.active
    assert ws["A1"].value == "Sensor"
    assert ws["A2"].value == "PT-101"
    assert ws["B2"].value == 42
    assert "References" in wb.sheetnames


def test_pdf_writer_creates_valid_pdf():
    result = PdfWriter().create_pdf(
        task_id="phase7-test",
        title="Compliance Report",
        content="This report confirms compliance with SOP-42.",
        table={"headers": ["Check", "Result"], "rows": [["Torque", "Pass"]]},
    )
    assert result["status"] == "ok"
    assert result["filename"].endswith(".pdf")
    assert result["file_hash"] == sha256_file(result["path"])
    with open(result["path"], "rb") as f:
        header = f.read(5)
    assert header == b"%PDF-"


def test_csv_and_json_writers_round_trip():
    csv_result = CsvWriter().create_csv(
        task_id="phase7-test", title="Export", headers=["a", "b"], rows=[[1, 2], [3, 4]]
    )
    assert csv_result["status"] == "ok"
    with open(csv_result["path"], encoding="utf-8") as f:
        content = f.read()
    assert "a,b" in content and "1,2" in content

    json_result = JsonWriter().create_json(task_id="phase7-test", title="Export", data={"key": "value"})
    assert json_result["status"] == "ok"
    with open(json_result["path"], encoding="utf-8") as f:
        loaded = json.load(f)
    assert loaded == {"key": "value"}


def test_docx_references_section_accepts_incoming_phase6_citation_shape():
    """The Phase 6 payload uses {source, text, page} instead of {doc_id, text} —
    both must render into the same 'References' section without error."""
    result = DocxWriter().create_docx(
        task_id="phase7-test",
        title="Valve Report",
        content="Findings below.",
        citations=[
            {"source": "SOP-14.pdf", "text": "depressurize to 450 PSI", "page": 4},
            {"doc_id": "sop-1", "text": "legacy chunk shape still supported"},
        ],
    )
    from docx import Document
    doc = Document(result["path"])
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "References" in full_text
    assert "SOP-14.pdf" in full_text and "p. 4" in full_text and "depressurize to 450 PSI" in full_text
    assert "sop-1" in full_text and "legacy chunk shape still supported" in full_text


def test_pdf_references_section_accepts_incoming_phase6_citation_shape():
    result = PdfWriter().create_pdf(
        task_id="phase7-test",
        title="Valve Report",
        content="Findings below.",
        citations=[{"source": "SOP-14.pdf", "text": "depressurize to 450 PSI", "page": 4}],
    )
    assert result["status"] == "ok"
    with open(result["path"], "rb") as f:
        assert f.read(5) == b"%PDF-"


def test_generate_docx_tool_accepts_citations_kwarg():
    from tools.registry import tool_registry
    result = tool_registry.execute_tool(
        "generate_docx", user_role="operator",
        arguments={
            "task_id": "phase7-test", "title": "Cited Report", "content": "Body",
            "citations": [{"source": "SOP-14.pdf", "text": "depressurize to 450 PSI", "page": 4}],
        },
    )
    assert result["status"] == "ok"
    from docx import Document
    doc = Document(result["path"])
    assert any("SOP-14.pdf" in p.text for p in doc.paragraphs)


def test_artifact_hash_changes_if_content_differs():
    r1 = DocxWriter().create_docx(task_id="phase7-test", title="A", content="Hello")
    r2 = DocxWriter().create_docx(task_id="phase7-test", title="A", content="Goodbye")
    assert r1["file_hash"] != r2["file_hash"]


# ── Tool registry integration ────────────────────────────────────────────────

def test_artifact_tools_registered():
    from tools.registry import tool_registry, ARTIFACT_TOOL_NAMES
    for name in ("generate_docx", "generate_xlsx", "generate_pdf", "generate_csv", "generate_json"):
        assert name in tool_registry.list_tools()
        assert name in ARTIFACT_TOOL_NAMES


def test_generate_docx_tool_callable_via_registry():
    from tools.registry import tool_registry
    result = tool_registry.execute_tool(
        "generate_docx", user_role="operator",
        arguments={"task_id": "phase7-test", "title": "Via Registry", "content": "Body text"},
    )
    assert result["status"] == "ok"
    assert os.path.isfile(result["path"])


def test_generate_docx_tool_permission_denied_for_viewer():
    from tools.registry import tool_registry
    with pytest.raises(PermissionError):
        tool_registry.execute_tool(
            "generate_docx", user_role="viewer",
            arguments={"task_id": "phase7-test", "title": "X", "content": "Y"},
        )


# ── Planner wiring ────────────────────────────────────────────────────────────

def test_determine_required_outputs_detects_deliverable_types():
    p = Planner()
    assert p.determine_required_outputs("Please generate a report and export as a document") == "docx"
    assert p.determine_required_outputs("Give me this data as an excel spreadsheet") == "xlsx"
    assert p.determine_required_outputs("Export the results as a pdf") == "pdf"
    assert p.determine_required_outputs("Give me a csv of the readings") == "csv"
    assert p.determine_required_outputs("What's the capital of France?") is None


def test_planner_appends_artifact_step_when_llm_plan_omits_it(monkeypatch):
    from models.registry import registry

    def mock_execute_prompt(model_id, prompt):
        # LLM forgets the deliverable rule and returns a plain reasoning plan.
        return json.dumps({"steps": [
            {"step_id": "step_1", "action": "Summarize findings", "tool": "direct_llm", "depends_on": [], "params": {}}
        ]}), 10.0

    monkeypatch.setattr(registry, "execute_prompt", mock_execute_prompt)
    p = Planner()
    plan = p.generate_plan("task-x", "Summarize the findings and generate a report document")
    assert any(s.tool == "generate_docx" for s in plan.steps)
    artifact_step = next(s for s in plan.steps if s.tool == "generate_docx")
    assert "step_1" in artifact_step.depends_on


def test_planner_does_not_append_artifact_step_when_not_requested(monkeypatch):
    from models.registry import registry

    def mock_execute_prompt(model_id, prompt):
        return json.dumps({"steps": [
            {"step_id": "step_1", "action": "Explain gravity", "tool": "direct_llm", "depends_on": [], "params": {}}
        ]}), 10.0

    monkeypatch.setattr(registry, "execute_prompt", mock_execute_prompt)
    p = Planner()
    plan = p.generate_plan("task-y", "Explain gravity")
    assert not any(s.tool in ("generate_docx", "generate_xlsx", "generate_pdf") for s in plan.steps)


# ── Full agent-loop integration ──────────────────────────────────────────────

def test_agent_produces_downloadable_artifact_end_to_end(monkeypatch):
    """DoD: the agent ends a report-generation task with a real, downloadable,
    hash-verifiable DOCX artifact, recorded in the `artifacts` table."""
    from models.registry import registry

    def mock_execute_prompt(model_id, prompt):
        if "You are a task planner." in prompt:
            plan = {"steps": [
                {"step_id": "s1", "action": "Draft findings", "tool": "direct_llm", "depends_on": [], "params": {"prompt": "Draft findings"}},
                {"step_id": "s2", "action": "Generate report", "tool": "generate_docx", "depends_on": ["s1"],
                 "params": {"title": "Inspection Report"}},
            ]}
            return json.dumps(plan), 50.0
        return "The equipment passed inspection with no defects found.", 50.0

    monkeypatch.setattr(registry, "execute_prompt", mock_execute_prompt)

    response = client.post("/api/agent", json={"prompt": "Draft findings and generate a report document"})
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "COMPLETED"

    artifact_events = [e for e in data["events"] if e["type"] == "artifact_created"]
    assert len(artifact_events) == 1
    artifact_id = artifact_events[0]["artifact_id"]
    assert artifact_events[0]["filename"].endswith(".docx")

    # Persisted to the artifacts table
    db = TestingSessionLocal()
    artifact = db.query(Artifact).filter(Artifact.id == artifact_id).first()
    db.close()
    assert artifact is not None
    assert artifact.task_id == data["task_id"]
    assert artifact.content_type.endswith("wordprocessingml.document")

    # Downloadable via the API, with the correct bytes and content-type
    dl = client.get(f"/api/artifacts/{artifact_id}/download")
    assert dl.status_code == 200
    assert dl.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert len(dl.content) > 0

    # Listed under its task
    listing = client.get(f"/api/artifacts?task_id={data['task_id']}")
    assert listing.status_code == 200
    assert any(a["id"] == artifact_id for a in listing.json())

    # Hash verifies clean
    verify = client.get(f"/api/artifacts/{artifact_id}/verify")
    assert verify.status_code == 200
    assert verify.json()["verified"] is True


def test_artifact_download_detects_tampering():
    result = DocxWriter().create_docx(task_id="phase7-test", title="Tamper Test", content="Original content")
    db = TestingSessionLocal()
    artifact = Artifact(
        id=result["id"], task_id="phase7-test", filename=result["filename"],
        content_type=result["content_type"], file_hash=result["file_hash"],
    )
    db.add(artifact)
    db.commit()
    artifact_id = artifact.id
    db.close()

    # Corrupt the file on disk after the hash was recorded
    with open(result["path"], "ab") as f:
        f.write(b"TAMPERED")

    dl = client.get(f"/api/artifacts/{artifact_id}/download")
    assert dl.status_code == 409

    verify = client.get(f"/api/artifacts/{artifact_id}/verify")
    assert verify.json()["verified"] is False


def test_artifact_download_missing_file_returns_410():
    db = TestingSessionLocal()
    artifact = Artifact(
        task_id="phase7-test", filename="ghost.docx",
        content_type="application/octet-stream", file_hash="deadbeef",
    )
    db.add(artifact)
    db.commit()
    db.refresh(artifact)
    artifact_id = artifact.id
    db.close()

    dl = client.get(f"/api/artifacts/{artifact_id}/download")
    assert dl.status_code == 410


def test_artifact_download_unknown_id_returns_404():
    dl = client.get("/api/artifacts/does-not-exist/download")
    assert dl.status_code == 404
