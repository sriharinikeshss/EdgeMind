"""
XLSX Artifact Writer — Phase 7 implementation (M5).

Uses openpyxl to build a workbook with a header row, data rows, and an
optional citations sheet.
"""
from __future__ import annotations
import uuid
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from artifacts.storage import artifact_path, sha256_file
from artifacts.citations import normalize_citation

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class XlsxWriter:
    """Writes task outputs to real .xlsx files, hashed and stored on disk."""

    def add_sheet_table(self, ws, headers: list[str], rows: list[list[Any]]) -> None:
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        for col_cells in ws.columns:
            width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
            ws.column_dimensions[col_cells[0].column_letter].width = min(width + 2, 60)

    def add_citations_sheet(self, wb: Workbook, citations: list[dict | str]) -> None:
        """Append a 'References' sheet. Accepts either citation shape —
        {"doc_id"/"source": ..., "text": ..., "page": ...} — via normalize_citation()."""
        ws = wb.create_sheet("References")
        ws.append(["#", "Source", "Page", "Text"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for i, citation in enumerate(citations, start=1):
            c = normalize_citation(citation)
            ws.append([i, c["source"], c["page"], c["text"]])

    def create_xlsx(
        self,
        task_id: str,
        title: str,
        headers: list[str],
        rows: list[list[Any]],
        citations: list[dict | str] | None = None,
        filename: str | None = None,
    ) -> dict:
        """Build an XLSX artifact and persist it to disk."""
        artifact_id = str(uuid.uuid4())
        filename = filename or f"{_slugify(title)}.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = title[:31] if title else "Sheet1"  # Excel sheet-name length limit
        self.add_sheet_table(ws, headers, rows)

        if citations:
            self.add_citations_sheet(wb, citations)

        path = artifact_path(task_id, artifact_id, filename)
        wb.save(path)
        file_hash = self.calculate_artifact_hash(path)

        return {
            "status": "ok",
            "id": artifact_id,
            "task_id": task_id,
            "filename": filename,
            "content_type": CONTENT_TYPE,
            "file_hash": file_hash,
            "path": path,
        }

    def calculate_artifact_hash(self, path: str) -> str:
        return sha256_file(path)


def _slugify(title: str) -> str:
    slug = "".join(c if c.isalnum() else "_" for c in title.strip()) or "artifact"
    return slug[:50]
