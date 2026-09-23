"""
DOCX Artifact Writer — Phase 7 implementation (M5).

Uses python-docx to build real documents with headings, paragraphs,
tables, and an appended citations block sourced from RAG metadata.

Supports markdown-style content parsing:
  - Lines starting with # or ## → Heading 1 / Heading 2
  - Lines starting with - [ ] or * [ ] → Checkbox checklist items (☐)
  - Lines starting with - or * → Bullet list items
  - Everything else → Regular paragraph text
"""
from __future__ import annotations
import uuid
from typing import Any

from docx import Document
from docx.shared import Pt

from artifacts.storage import artifact_path, sha256_file
from artifacts.citations import format_citation_line

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class DocxWriter:
    """Writes task outputs to real .docx files, hashed and stored on disk."""

    def add_heading(self, doc: Document, text: str, level: int = 1) -> None:
        doc.add_heading(text, level=level)

    def add_table(self, doc: Document, headers: list[str], rows: list[list[Any]]) -> None:
        table = doc.add_table(rows=1, cols=len(headers))
        try:
            table.style = "Light Grid Accent 1"
        except KeyError:
            pass  # style not available in a minimal template — fall back to default grid
        header_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            header_cells[i].text = str(header)
        for row in rows:
            row_cells = table.add_row().cells
            for i, value in enumerate(row):
                row_cells[i].text = str(value)

    def add_citations(self, doc: Document, citations: list[dict | str]) -> None:
        """Append a 'References' section. Accepts either citation shape —
        {"doc_id"/"source": ..., "text": ..., "page": ...} — via normalize_citation()."""
        self.add_heading(doc, "References", level=2)
        for i, citation in enumerate(citations, start=1):
            doc.add_paragraph(format_citation_line(i, citation))

    def apply_template(self, doc: Document, template_path: str) -> Document:
        """Load a corporate .docx template as the base document instead of a blank one."""
        return Document(template_path)

    def _add_content_rich(self, doc: Document, content: str) -> None:
        """
        Parse content string with markdown-style formatting into structured DOCX elements.

        Supports:
          ## Heading 2 text        → doc.add_heading(level=2)
          # Heading 1 text         → doc.add_heading(level=1)
          - [ ] checklist item     → paragraph with ☐ prefix (List Bullet style)
          - [x] checked item       → paragraph with ☑ prefix (List Bullet style)
          - bullet item            → List Bullet style paragraph
          * bullet item            → List Bullet style paragraph
          **bold text**            → bold run paragraph
          normal text              → regular paragraph
        """
        # Support both literal \n in JSON strings and real newlines
        lines = content.replace("\\n", "\n").split("\n")

        for line in lines:
            stripped = line.strip()

            if not stripped:
                continue

            # Heading 2: ##
            if stripped.startswith("## "):
                doc.add_heading(stripped[3:].strip(), level=2)

            # Heading 1: #
            elif stripped.startswith("# "):
                doc.add_heading(stripped[2:].strip(), level=1)

            # Unchecked checkbox: - [ ] or * [ ]
            elif stripped.startswith(("- [ ]", "* [ ]")):
                text = stripped[5:].strip()
                p = doc.add_paragraph(style="List Bullet")
                run = p.add_run(f"\u2610  {text}")
                run.font.size = Pt(11)

            # Checked checkbox: - [x] or * [x]
            elif stripped.startswith(("- [x]", "* [x]", "- [X]", "* [X]")):
                text = stripped[5:].strip()
                p = doc.add_paragraph(style="List Bullet")
                run = p.add_run(f"\u2611  {text}")
                run.font.size = Pt(11)

            # Bullet point: - or *
            elif stripped.startswith("- ") or stripped.startswith("* "):
                text = stripped[2:].strip()
                doc.add_paragraph(text, style="List Bullet")

            # Bold: **text**
            elif stripped.startswith("**") and stripped.endswith("**") and len(stripped) > 4:
                p = doc.add_paragraph()
                run = p.add_run(stripped[2:-2])
                run.bold = True

            # Regular paragraph
            else:
                doc.add_paragraph(stripped)

    def create_docx(
        self,
        task_id: str,
        title: str,
        content: str,
        table: dict | None = None,
        citations: list[dict | str] | None = None,
        template_path: str | None = None,
        filename: str | None = None,
    ) -> dict:
        """
        Build a DOCX artifact and persist it to disk.

        `content` supports markdown-style formatting (##, -, - [ ], **bold**).
        `table`, if given, is {"headers": [...], "rows": [[...], ...]}.
        Returns artifact metadata: id, filename, content_type, file_hash, path.
        """
        artifact_id = str(uuid.uuid4())
        filename = filename or f"{_slugify(title)}.docx"

        doc = self.apply_template(None, template_path) if template_path else Document()
        self.add_heading(doc, title, level=1)

        if content and content.strip():
            self._add_content_rich(doc, content)
        else:
            doc.add_paragraph("No content was provided for this document.")

        if table and table.get("headers"):
            self.add_table(doc, table["headers"], table.get("rows", []))

        if citations:
            self.add_citations(doc, citations)

        path = artifact_path(task_id, artifact_id, filename)
        doc.save(path)
        file_hash = self.calculate_artifact_hash(path)

        return {
            "status": "ok",
            "id": artifact_id,
            "task_id": task_id,
            "filename": filename,
            "content_type": CONTENT_TYPE,
            "file_hash": file_hash,
            "path": path,
        }

    def save_artifact(self, doc: Document, path: str) -> str:
        """Save an in-memory Document to `path` and return its SHA-256 hash."""
        doc.save(path)
        return self.calculate_artifact_hash(path)

    def calculate_artifact_hash(self, path: str) -> str:
        return sha256_file(path)


def _slugify(title: str) -> str:
    slug = "".join(c if c.isalnum() else "_" for c in title.strip()) or "artifact"
    return slug[:50]
